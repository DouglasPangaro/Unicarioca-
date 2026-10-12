import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

from app.core.erros import Invalido
from app.domain.modelos import Lado, q2
from app.domain.tickers import normalizar_ticker

# campo interno -> cabeçalho da planilha "Negociação" da Área do Investidor B3
COLUNAS = {
    "data": "Data do Negócio",
    "lado": "Tipo de Movimentação",
    "mercado": "Mercado",
    "instituicao": "Instituição",
    "ticker": "Código de Negociação",
    "quantidade": "Quantidade",
    "preco": "Preço",
}
MERCADOS_ACEITOS = {"mercado a vista", "mercado fracionario"}
LADOS = {"compra": Lado.COMPRA, "venda": Lado.VENDA}


def normalizar_texto(valor: object) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return " ".join(texto.lower().split())


@dataclass(frozen=True)
class LinhaB3:
    numero: int
    data: date
    lado: Lado
    instituicao: str
    ticker: str
    quantidade: int
    preco: Decimal


@dataclass
class LeituraB3:
    linhas: list[LinhaB3] = field(default_factory=list)
    ignoradas: list[tuple[int, str]] = field(default_factory=list)


def _data(valor: object) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return datetime.strptime(str(valor).strip(), "%d/%m/%Y").date()


def _decimal(valor: object) -> Decimal:
    if isinstance(valor, int | float):
        return Decimal(str(valor))
    texto = str(valor).replace("R$", "").strip()
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    return Decimal(texto)


def _quantidade(valor: object) -> int:
    numero = _decimal(valor)
    if numero <= 0 or numero != numero.to_integral_value():
        raise ValueError(f"quantidade inválida: {valor}")
    return int(numero)


def _preco(valor: object) -> Decimal:
    preco = q2(_decimal(valor))
    if preco <= 0:
        raise ValueError(f"preço inválido: {valor}")
    return preco


def _instituicao(valor: object) -> str:
    nome = " ".join(str(valor or "").split())
    if not nome or len(nome) > 100:
        raise ValueError("instituição ausente ou com mais de 100 caracteres")
    return nome


def ler_planilha(conteudo: bytes, max_linhas: int) -> LeituraB3:
    try:
        planilha = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:  # qualquer falha de leitura de arquivo enviado pelo usuário é entrada inválida
        raise Invalido("O arquivo não é uma planilha .xlsx válida.") from erro
    try:
        linhas = planilha.worksheets[0].iter_rows(values_only=True)
        nomes = {normalizar_texto(original): campo for campo, original in COLUNAS.items()}
        cabecalho = [normalizar_texto(c) for c in (next(linhas, None) or ())]
        indices = {nomes[n]: i for i, n in enumerate(cabecalho) if n in nomes}
        ausentes = [COLUNAS[c] for c in COLUNAS if c not in indices]
        if ausentes:
            raise Invalido("A planilha não tem as colunas esperadas da B3.", colunas_ausentes=ausentes)

        leitura = LeituraB3()
        for numero, linha in enumerate(linhas, start=2):
            if all(v in (None, "") for v in linha):
                continue
            if numero - 1 > max_linhas:
                raise Invalido(f"A planilha excede o limite de {max_linhas} linhas.")
            valores = {c: (linha[i] if i < len(linha) else None) for c, i in indices.items()}
            if normalizar_texto(valores["mercado"]) not in MERCADOS_ACEITOS:
                leitura.ignoradas.append((numero, f"Mercado não suportado: {valores['mercado']}"))
                continue
            lado = LADOS.get(normalizar_texto(valores["lado"]))
            if lado is None:
                leitura.ignoradas.append((numero, f"Tipo de movimentação desconhecido: {valores['lado']}"))
                continue
            try:
                leitura.linhas.append(
                    LinhaB3(
                        numero=numero,
                        data=_data(valores["data"]),
                        lado=lado,
                        instituicao=_instituicao(valores["instituicao"]),
                        ticker=normalizar_ticker(str(valores["ticker"] or "")),
                        quantidade=_quantidade(valores["quantidade"]),
                        preco=_preco(valores["preco"]),
                    )
                )
            except (ValueError, TypeError, InvalidOperation) as erro:
                leitura.ignoradas.append((numero, f"Linha inválida: {erro}"))
        return leitura
    finally:
        planilha.close()
